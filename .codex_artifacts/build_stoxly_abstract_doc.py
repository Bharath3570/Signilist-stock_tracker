from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUT_PATH = Path(r"C:\Users\Varuna\WebstormProjects\stocks_app\Stoxly_Abstract.docx")

def main():
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    title = doc.add_paragraph('ABSTRACT')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.runs[0]
    r.bold = True
    r.font.size = Pt(16)

    doc.add_paragraph('')

    abstract_paras = [
        "Stoxly is a real-time stock market web application designed to help retail investors track equities, monitor personal watchlists, and receive timely market updates in a single, easy-to-use platform. Traditional market tracking requires users to switch between multiple sources for quotes, company information, charts, and news, which increases effort and delays decision-making. Stoxly addresses this by combining stock discovery, symbol-level views, personalized tracking, and automated notifications into an integrated end-to-end system.",
        "The application provides a secure authentication workflow for account creation and sign-in, and then presents a dashboard-driven user experience. Users can search for stocks, view a dedicated stock details page for each symbol, and manage a watchlist for fast access to frequently monitored companies. To reduce information overload, the system integrates market news retrieval for watched symbols and supports generating short summaries so users can quickly understand key developments without reading multiple full articles.",
        "A key contribution of Stoxly is its background automation pipeline for alerts and engagement. Users can configure price alerts (greater-than or less-than conditions) for selected symbols. A scheduled background workflow periodically checks current prices and evaluates alert conditions. When an alert is triggered, Stoxly records the trigger in alert history, creates a user notification, and dispatches an email message to ensure the user receives the update even when not actively using the application. This automation is implemented using Inngest for reliable event-driven and cron-based workflows, while email delivery is handled through Nodemailer.",
        "The system is implemented using Next.js with modern full-stack patterns, including server-side actions for backend operations and a component-driven UI for maintainable frontend development. Data persistence is provided by MongoDB, with Mongoose models managing collections such as watchlists, alerts, alert history, and notifications. The solution also integrates external market data and news sources to fetch quotes and relevant information in near real time.",
        "Stoxly is deployed on Vercel, enabling production-ready hosting with scalable delivery. Overall, the project demonstrates a practical stock tracking platform that combines real-time visibility, personalization, automated monitoring, and multi-channel communication. Future enhancements can include richer analytics, push notifications, improved caching for faster performance, and expanded coverage across additional markets and asset classes.",
    ]

    for text in abstract_paras:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15

    doc.save(str(OUT_PATH))
    print(str(OUT_PATH))

if __name__ == '__main__':
    main()
