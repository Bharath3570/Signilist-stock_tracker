from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

PROJECT_ROOT = Path(r"C:\Users\Varuna\WebstormProjects\stocks_app")
OUT_PATH = PROJECT_ROOT / "Stoxly_Code_Appendix.docx"

FILES = [
    r"app/(auth)/sign-up/page.tsx",
    r"app/(auth)/sign-in/page.tsx",
    r"app/(root)/watchlist/page.tsx",
    r"app/(root)/stocks/[symbol]/page.tsx",
    r"app/(root)/alerts/page.tsx",
    r"app/(root)/notifications/page.tsx",
    r"app/api/inngest/route.ts",
    r"lib/better-auth/auth.ts",
    r"lib/actions/auth.actions.ts",
    r"database/mongoose.ts",
    r"database/models/watchlist.model.ts",
    r"database/models/alert.model.ts",
    r"database/models/notification.model.ts",
    # optional
    r"database/models/alert-history.model.ts",
    r"database/models/insight.model.ts",
    r"lib/actions/finnhub.actions.ts",
    r"lib/actions/watchlist.actions.ts",
    r"lib/actions/alert.actions.ts",
    r"lib/actions/notification.actions.ts",
    r"lib/inngest/client.ts",
    r"lib/inngest/functions.ts",
    r"lib/nodemailer/index.ts",
    r"components/ui/SearchCommand.tsx",
    r"components/ui/TradingViewWidget.tsx",
    r"components/ui/CreateAlertModal.tsx",
    r"components/ui/WatchlistTable.tsx",
]


def set_cell_shading(cell, fill: str = "F3F4F6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_monospace(run, size_pt: int = 9):
    run.font.name = "Consolas"
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:ascii"), "Consolas")
    rfonts.set(qn("w:hAnsi"), "Consolas")
    rfonts.set(qn("w:eastAsia"), "Consolas")
    rfonts.set(qn("w:cs"), "Consolas")


def add_code_block(doc: Document, code: str):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    set_cell_shading(cell)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    # Preserve newlines exactly
    lines = code.splitlines()
    if not lines:
        lines = [""]

    for i, line in enumerate(lines):
        run = p.add_run(line)
        set_run_monospace(run)
        if i != len(lines) - 1:
            run.add_break()


def main():
    doc = Document()

    # Basic page setup
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    # Title page (placeholders as requested)
    title = doc.add_paragraph("STOXLY")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(22)

    subtitle = doc.add_paragraph("Real-Time Stock Market Web Application")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.size = Pt(14)

    doc.add_paragraph("")
    p = doc.add_paragraph("Project Report Code Appendix")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)

    doc.add_paragraph("")
    for label in [
        "Student Name:",
        "Register Number / Roll No:",
        "Project Guide:",
        "College / Department:",
        "Academic Year:",
        "Deployed URL:",
        "GitHub Repository URL:",
    ]:
        line = doc.add_paragraph(f"{label} ________________________________")

    doc.add_page_break()

    # Contents
    h = doc.add_paragraph("Contents")
    h.runs[0].bold = True
    h.runs[0].font.size = Pt(16)
    doc.add_paragraph("(This appendix includes key source files used in the Stoxly project.)")
    for rel in FILES:
        if (PROJECT_ROOT / rel).exists():
            doc.add_paragraph(rel, style="List Bullet")
    doc.add_page_break()

    # Code sections
    for rel in FILES:
        path = PROJECT_ROOT / rel
        if not path.exists():
            continue

        doc.add_paragraph(rel).runs[0].bold = True
        doc.paragraphs[-1].runs[0].font.size = Pt(13)

        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = path.read_text(encoding="utf-8", errors="replace")

        add_code_block(doc, text)
        doc.add_paragraph("")

    doc.save(str(OUT_PATH))
    print(str(OUT_PATH))


if __name__ == "__main__":
    main()
